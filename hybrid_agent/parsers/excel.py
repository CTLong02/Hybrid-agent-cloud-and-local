"""Excel parser. Header row defines columns; one task per data row.

Required column: `id`. Recognized: title, description, depends_on, tags,
complexity, target_files (or files), acceptance_criteria (or acceptance).
List-typed columns accept comma- or semicolon-separated values.
"""

from __future__ import annotations

from pathlib import Path

from ..models import TaskSpec
from .base import build_spec, register


class ExcelParser:
    extensions: tuple[str, ...] = (".xlsx", ".xlsm")

    def parse(self, path: Path) -> list[TaskSpec]:
        from openpyxl import load_workbook  # lazy import — heavy dep

        wb = load_workbook(filename=str(path), data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h or "").strip() for h in rows[0]]
        out: list[TaskSpec] = []
        for row in rows[1:]:
            d = {h: (v if v is not None else "") for h, v in zip(headers, row, strict=False) if h}
            if not str(d.get("id") or "").strip():
                continue
            out.append(build_spec(d))
        return out


register(ExcelParser())
